import docx        # for DOCX
import openpyxl    # for XLSX
import pptx        # for PPTX
import pytesseract # for Images (OCR)
from PIL import Image
import io
import pymupdf
async def extract_text_universally(content_bytes, file_type):
    text = ""
    
    # 1. PDF 
    if file_type == "application/pdf":
        pdf=pymupdf.open(stream=content_bytes, filetype="pdf")
        text = "".join([page.get_text() for page in pdf])
    
    # 2. DOCX 
    elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(io.BytesIO(content_bytes))
        text = "\n".join([para.text for para in doc.paragraphs])

    # 3. XLSX 
    elif file_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"

    # 4. PPTX
    elif file_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        pres = pptx.Presentation(io.BytesIO(content_bytes))
        for slide in pres.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"

    # 5. Images (OCR)
    elif file_type in ["image/jpeg", "image/png", "image/tiff"]:
        image = Image.open(io.BytesIO(content_bytes))
        text = pytesseract.image_to_string(image)

    # 6. TXT
    elif file_type == "text/plain":
        text = content_bytes.decode("utf-8", errors="ignore")

    return text